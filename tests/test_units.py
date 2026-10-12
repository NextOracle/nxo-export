"""Unit tests for the seams of the OOP design: styles catalogue, link resolution, HTML assembly, OOXML passes."""
from collections import Counter
from pathlib import Path

from docx import Document
from docx.oxml.ns import qn
from PIL import Image

from nxo_export.config import DocxCfg
from nxo_export.domain.site import Site, TocNode
from nxo_export.exporters.docx.html import HtmlAssembler, LinkResolver
from nxo_export.exporters.docx.ooxml import PassContext, PostProcessor
from nxo_export.exporters.docx.ooxml import xml as ox
from nxo_export.exporters.docx.ooxml.passes import DocumentPass, FigurePass
from nxo_export.exporters.docx.ooxml.passes.base import MIN_CELL_CONTENT, content_width
from nxo_export.exporters.docx.styles import WritersideStyleCatalogue
from nxo_export.exporters.docx.styles.specs import ParagraphStyle, TableStyle
from nxo_export.exporters.docx.theme import Theme
from nxo_export.exporters.docx.topic import TopicHtml

# ------------------------------------------------------------------ styles


def test_catalogue_declares_every_pandoc_style_once():
    specs = WritersideStyleCatalogue(Theme(), DocxCfg(), 9000).styles()
    names = [s.name for s in specs]
    assert len(names) == len(set(names))
    assert names[0] == "Normal"
    for needed in ("Heading 1", "Heading 6", "Source Code", "Text Diagram", "Hyperlink", "toc 1", "Callout Tip"):
        assert needed in names
    assert sum(isinstance(s, TableStyle) for s in specs) == 1
    toc1 = next(s for s in specs if s.name == "toc 1")
    assert isinstance(toc1, ParagraphStyle) and toc1.right_tab_twips == 9000


def test_catalogue_honours_font_overrides():
    specs = WritersideStyleCatalogue(Theme(), DocxCfg(font="Arial", code_font="Consolas"), 9000).styles()
    by_name = {s.name: s for s in specs}
    assert by_name["Normal"].font == "Arial" and by_name["Source Code"].font == "Consolas"


# ------------------------------------------------------------------ html


def _topic(url: str, slug_ids=(), level: int = 0, html: str = "<h1>x</h1>") -> TopicHtml:
    node = TocNode(id=url, title=url.title(), url=url, level=level)
    return TopicHtml(node, html, 0, anchor=f"{node.slug}--top", ids=frozenset(slug_ids))


def test_link_resolver_internal_external_and_plain():
    a = _topic("a.html")
    b = _topic("b.html", slug_ids={"b--sec"})
    topics = {"a.html": a, "b.html": b}
    html = '<a href="nxo:b.html#sec">B</a><a href="nxo:c.html">C</a><a href="nxo:#x">self</a>'
    out = LinkResolver(topics, None).resolve(html, current=a, internal=lambda u: u in topics)
    assert 'href="#b--sec"' in out and "<span>C</span>" in out and 'href="#a--top"' in out
    ext = LinkResolver(topics, "https://docs/").resolve(html, current=a, internal=lambda u: u == "a.html")
    assert 'href="https://docs/b.html#sec"' in ext and 'href="https://docs/c.html"' in ext


def test_combined_document_groups_and_shifts_headings():
    t = _topic("t.html", level=1, html='<h1 id="t">T</h1>')
    toc = (TocNode("g", "Group", None, 0, (t.node, TocNode("x", "X", "x.html", 1))),)
    html = HtmlAssembler({"t.html": t}, None).combined(Site(Path("."), toc, "P"), "Doc")
    assert "<title>Doc</title>" in html
    assert '<h1 id="nxoa1">Group</h1>' in html and '<h2 id="nxoa2">T</h2>' in html and "x.html" not in html


# ------------------------------------------------------------------ ooxml


def test_content_width_subtracts_cell_margins():
    body = ox.make("body")
    tbl = ox.make("tbl")
    tblpr = ox.make("tblPr")
    ox.margins(tblpr, "tblCellMar", 0, 150, 0, 150)
    tbl.append(tblpr)
    tc = ox.make("tc")
    tcpr = ox.make("tcPr")
    tcpr.append(ox.make("tcW", w=3000, type="dxa"))
    tc.append(tcpr)
    tbl.append(tc)
    body.append(tbl)
    assert content_width(body, 9000) == 9000
    assert content_width(tc, 9000) == 2700
    tcpr.find(qn("w:tcW")).set(qn("w:w"), "250")
    assert content_width(tc, 9000) == MIN_CELL_CONTENT


def test_figure_pass_clamps_and_drops_source_path(tmp_path):
    png = tmp_path / "big.png"
    Image.new("RGB", (4000, 400), "white").save(png)
    doc = Document()
    doc.add_paragraph().add_run().add_picture(str(png))
    for c in doc.element.body.iter("{http://schemas.openxmlformats.org/drawingml/2006/picture}cNvPr"):
        c.set("descr", "/private/tmp/secret/big.png")
    stats: Counter[str] = Counter()
    FigurePass().run(doc, PassContext(theme=Theme(), text_width=9000), stats)
    assert stats["figures"] == 1
    xml = doc.element.xml
    assert "/private/tmp" not in xml
    ext = doc.element.body.find(".//{http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing}extent")
    assert int(ext.get("cx")) <= 9000 * 635


def test_post_processor_runs_only_applicable_passes(tmp_path):
    seen: list[str] = []

    class Probe(DocumentPass):
        def __init__(self, name: str, combined_only: bool):
            self.name, self.combined_only = name, combined_only

        def applies(self, ctx):
            return ctx.combined or not self.combined_only

        def run(self, doc, ctx, stats):
            seen.append(self.name)
            stats[self.name] += 1

    out = tmp_path / "d.docx"
    Document().save(str(out))
    stats = PostProcessor([Probe("always", False), Probe("combined", True)]).run(
        out, PassContext(theme=Theme(), text_width=9000))
    assert seen == ["always"] and stats == Counter(always=1)
