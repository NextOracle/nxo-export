"""End-to-end export of a real Writerside project (runs the builder in Docker + Chromium + Pandoc).

    WS_E2E_PROJECT=/path/to/Writerside WS_E2E_INSTANCE=nsk pytest -m e2e
"""

import io
import os
import zipfile
from pathlib import Path

import pytest
from click.testing import CliRunner
from docx import Document
from PIL import Image, ImageStat
from pypdf import PdfReader

from nxo_ws_export.cli import main

pytestmark = pytest.mark.e2e

PROJECT = os.environ.get("WS_E2E_PROJECT")
INSTANCE = os.environ.get("WS_E2E_INSTANCE")


@pytest.mark.skipif(not PROJECT, reason="WS_E2E_PROJECT not set")
@pytest.mark.parametrize("theme", ["light", "dark"])
def test_export_real_project(tmp_path: Path, theme: str):
    args = ["export", "-p", PROJECT, "-f", "pdf", "-t", theme, "-m", "both", "-o", str(tmp_path)]
    if theme == "light":
        args += ["-f", "docx"]
    if INSTANCE:
        args += ["-i", INSTANCE]
    res = CliRunner().invoke(main, args, catch_exceptions=False)
    assert res.exit_code == 0, res.output

    inst_dir = next(p for p in tmp_path.iterdir() if p.is_dir())
    pdf = PdfReader(inst_dir / f"{inst_dir.name}-{theme}.pdf")
    topics = sorted((inst_dir / f"pdf-{theme}").glob("*.pdf"))
    assert topics and len(pdf.pages) > len(topics)
    assert pdf.outline, "combined PDF must have bookmarks"

    if theme == "light":
        doc = Document(str(inst_dir / f"{inst_dir.name}.docx"))
        assert any(p.style.name.startswith("Heading") for p in doc.paragraphs)
        assert len(list((inst_dir / "docx").glob("*.docx"))) == len(topics)
        assert doc.tables and all(t.style is not None for t in doc.tables)
        # DOCX is always light: every embedded figure (Mermaid included) must be a light image.
        with zipfile.ZipFile(inst_dir / f"{inst_dir.name}.docx") as z:
            for name in (n for n in z.namelist() if n.startswith("word/media/")):
                rgba = Image.open(io.BytesIO(z.read(name))).convert("RGBA")
                img = Image.alpha_composite(Image.new("RGBA", rgba.size, "white"), rgba).convert("L")
                assert ImageStat.Stat(img).mean[0] > 128, f"{name} looks dark"
